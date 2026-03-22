import pandas as pd
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import datetime
import os

class ExcelProcessorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel Data Integrator - Mandy Version (V2.6 Warning Mode)")
        self.root.geometry("1400x900")
        
        self.input_file_path = ""
        self.output_target_path = ""
        self.all_sheets = []
        self.sheet_checkboxes = {} 
        
        # 標準定義清單 (我們期望看到的欄位)
        self.column_names_list = [
            'Order Date', 'Sold-to Party', 'Ship-to Party', 'SAP SO#', 'SO Item',
            'Customer PO#', 'line no.', 'Part Number', 'Customer Part Number',
            'Qty Ordered', 'Unit Price', 'Required. Date', 'shipping schedule',
            'Remark', 'Tracking no.', '(訂單變更) 新增項次', '來源項次',
            'End Customer PN', 'BU', 'Interface', 'Flash / DRAM', 'ETD reason', 'ETD History'
        ]
        
        self.mapping = {
            'Order Date': 'Customer Reference Date', 'Sold-to Party': 'Sold-to Party',
            'Ship-to Party': 'Ship-to Party', 'SAP SO#': 'Sales Document',
            'SO Item': 'Sales Document Item', 'Customer PO#': 'Customer Reference',
            'line no.': 'Customer PO Item No', 'Part Number': 'Material',
            'Customer Part Number': 'Customer Material', 'Qty Ordered': 'Order Quantity',
            'Unit Price': 'Unit Price', 'Required. Date': 'CRD',
            'shipping schedule': 'ETD', 'Remark': 'Remark', 'Tracking no.':'Price Check',
            'BU': 'BU', 'Interface': 'Interface', 'Flash / DRAM': 'Flash / DRAM',
            'ETD reason': 'ETD Reason', 'ETD History': 'ETD History'
        }

        self.mandy_columns_list = [
            'Upload', 'Close', 'SO Type', 'Source SO No', 'Source SO Item No', 'Sales Document',
            'Sales Document Item', 'Sold-to Abbreviation', 'Ship-to Abbreviation', 'Customer Reference Date',
            'Sold-to Party', 'Ship-to Party', 'Customer Reference', 'Customer PO Item No', 'Material',
            'Module Material', 'Customer Material', 'Order Quantity', 'Unit Price', 'CRD', 'ETD', 'ETA',
            'Is Sample (Y)', 'Remark', 'Price Check', 'Purchase Order Type', 'Sample/Loan SFID', 'Is Gift Item (ZFTO)',
            'FOC Invoice Price', 'Masking Price', 'Masking Price Currency', 'End Customer PO#', 'End Customer PN',
            'Incoterm (Header)', 'Inco.Location (Header)', 'Incoterm (Item)', 'Inco.Location (Item)',
            'Special Procurement Application No', 'Titop DN', 'Tiptop DN Item', 'Unit Price Currency',
            'Original Unit Price', 'BU', 'Interface', 'Flash / DRAM', 'ETD Reason', 'Allocation Date',
            'ETD History', '訂單變更類型', '拆項來源項次 / 取消原因', '(ZDR) \nItem Description'
        ]

        self.tree_view_cols = ['Customer Reference Date', 'Sold-to Party', 'Sales Document', 'Material', 'Order Quantity', 'ETD', 'Remark', 'BU']

        self.setup_ui()

    def setup_ui(self):
        main_paned = ttk.Panedwindow(self.root, orient=tk.HORIZONTAL)
        main_paned.pack(fill=tk.BOTH, expand=True)
        left_frame = ttk.Frame(main_paned, padding=15)
        main_paned.add(left_frame, weight=1)

        ttk.Label(left_frame, text="Step 1: Input File", font=('Arial', 10, 'bold')).pack(anchor=tk.W)
        ttk.Button(left_frame, text="📁 Open Source Excel", command=self.open_file).pack(fill=tk.X, pady=5)
        self.input_label = ttk.Label(left_frame, text="No source selected", wraplength=350, foreground="gray")
        self.input_label.pack(pady=2, anchor=tk.W)

        ttk.Separator(left_frame, orient='horizontal').pack(fill='x', pady=10)
        ttk.Label(left_frame, text="Step 2: Output Target", font=('Arial', 10, 'bold')).pack(anchor=tk.W)
        ttk.Button(left_frame, text="🎯 Set Target File", command=self.set_output_target).pack(fill=tk.X, pady=5)
        self.output_label = ttk.Label(left_frame, text="Default: Same as Source", wraplength=350, foreground="gray")
        self.output_label.pack(pady=2, anchor=tk.W)

        ttk.Separator(left_frame, orient='horizontal').pack(fill='x', pady=10)
        ttk.Label(left_frame, text="Step 3: Search Sheets", font=('Arial', 10, 'bold')).pack(anchor=tk.W)
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", self.filter_sheets)
        ttk.Entry(left_frame, textvariable=self.search_var).pack(fill=tk.X, pady=5)

        btn_tool_frame = ttk.Frame(left_frame)
        btn_tool_frame.pack(fill=tk.X)
        ttk.Button(btn_tool_frame, text="All Visible", command=self.select_all_visible).pack(side=tk.LEFT, expand=True, fill=tk.X)
        ttk.Button(btn_tool_frame, text="None Visible", command=self.clear_all_visible).pack(side=tk.LEFT, expand=True, fill=tk.X)

        list_container = ttk.Frame(left_frame, relief="sunken", borderwidth=1)
        list_container.pack(fill=tk.BOTH, expand=True, pady=5)
        self.canvas = tk.Canvas(list_container, highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(list_container, orient="vertical", command=self.canvas.yview)
        self.scroll_frame = ttk.Frame(self.canvas)
        self.scroll_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.skip_hidden_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(left_frame, text="Skip Hidden Rows", variable=self.skip_hidden_var).pack(anchor=tk.W)

        ttk.Button(left_frame, text="🚀 START PROCESS", command=self.process_data).pack(fill=tk.X, ipady=10, pady=10)

        right_frame = ttk.Frame(main_paned, padding=15)
        main_paned.add(right_frame, weight=3)
        self.info_label = ttk.Label(right_frame, text="Ready", font=('Arial', 11, 'bold'))
        self.info_label.pack(anchor=tk.W, pady=(0, 10))
        self.tree = ttk.Treeview(right_frame, show="headings")
        self.tree.pack(fill=tk.BOTH, expand=True)

    def open_file(self):
        file = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file: return
        self.input_file_path = file
        self.output_target_path = file
        self.input_label.config(text=os.path.basename(file), foreground="blue")
        try:
            xl = pd.ExcelFile(file)
            self.all_sheets = xl.sheet_names
            for widget in self.scroll_frame.winfo_children(): widget.destroy()
            self.sheet_checkboxes = {}
            for sheet in self.all_sheets:
                var = tk.BooleanVar(value=False)
                chk = ttk.Checkbutton(self.scroll_frame, text=sheet, variable=var)
                self.sheet_checkboxes[sheet] = (var, chk)
            self.filter_sheets()
        except Exception as e: messagebox.showerror("Error", str(e))

    def filter_sheets(self, *args):
        s = self.search_var.get().strip().lower()
        for name, (var, chk) in self.sheet_checkboxes.items():
            if s in name.lower(): chk.pack(anchor=tk.W, pady=1, fill=tk.X)
            else: chk.pack_forget()
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def set_output_target(self):
        file = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if file:
            self.output_target_path = file
            self.output_label.config(text=os.path.basename(file), foreground="blue")

    def select_all_visible(self):
        s = self.search_var.get().strip().lower()
        for name, (var, chk) in self.sheet_checkboxes.items():
            if s in name.lower(): var.set(True)

    def clear_all_visible(self):
        s = self.search_var.get().strip().lower()
        for name, (var, chk) in self.sheet_checkboxes.items():
            if s in name.lower(): var.set(False)

    def process_data(self):
        selected_names = [n for n, (v, c) in self.sheet_checkboxes.items() if v.get()]
        if not selected_names:
            messagebox.showwarning("Warning", "Please select at least one sheet!")
            return

        # --- 1. 預先檢查所有 Sheet 的欄位 ---
        warnings = []
        # 標準化我們的預期清單 (去空白換行)
        std_list = [str(c).strip().replace('\n', ' ') for c in self.column_names_list]

        for s_name in selected_names:
            try:
                # 只讀取第一列來檢查標題，速度最快
                df_cols = pd.read_excel(self.input_file_path, sheet_name=s_name, nrows=0).columns
                clean_df_cols = [str(c).strip().replace('\n', ' ') for c in df_cols]
                
                missing = [orig for orig, std in zip(self.column_names_list, std_list) if std not in clean_df_cols]
                if missing:
                    warnings.append(f"Sheet [{s_name}] 缺少欄位: {', '.join(missing)}")
            except:
                warnings.append(f"Sheet [{s_name}] 無法讀取欄位。")

        # --- 2. 顯示 Warning 並詢問是否繼續 ---
        if warnings:
            warn_msg = "發現以下欄位缺失或不符，遺漏的資料將以空值呈現：\n\n" + "\n".join(warnings) + "\n\n是否仍要繼續執行？"
            confirm = messagebox.askyesno("Field Warning", warn_msg)
            if not confirm:
                return # 使用者選擇取消

        # --- 3. 正式開始處理 ---
        self.info_label.config(text="Processing... Please wait.")
        self.root.update()
        processed_dfs = []
        
        try:
            wb_hidden = None
            if self.skip_hidden_var.get():
                wb_hidden = openpyxl.load_workbook(self.input_file_path, data_only=True)

            for s_name in selected_names:
                df_src = pd.read_excel(self.input_file_path, sheet_name=s_name, dtype=str)
                df_src.columns = [str(c).strip().replace('\n', ' ') for c in df_src.columns]
                
                # 建立對齊後的空 DF
                df_aligned = pd.DataFrame(columns=std_list)
                common_cols = [c for c in std_list if c in df_src.columns]
                df_aligned = pd.concat([df_aligned, df_src[common_cols]], axis=0, ignore_index=True)

                if wb_hidden and s_name in wb_hidden.sheetnames:
                    ws = wb_hidden[s_name]
                    if hasattr(ws, 'row_dimensions'):
                        hidden = [r for r, dim in ws.row_dimensions.items() if dim.hidden]
                        keep = [i for i in range(len(df_src)) if (i + 2) not in hidden]
                        df_aligned = df_aligned.iloc[keep]

                if 'line no.' in df_aligned.columns:
                    mask = df_aligned['line no.'].astype(str).str.strip().isin(['', 'nan', 'None'])
                    if mask.any(): df_aligned = df_aligned.iloc[:df_aligned.index.get_loc(mask.idxmax())]
                
                processed_dfs.append(df_aligned)

            if wb_hidden: wb_hidden.close()
            if not processed_dfs: return

            # 合併與映射
            final_raw = pd.concat(processed_dfs, ignore_index=True)
            clean_mapping = {str(k).strip().replace('\n', ' '): v for k, v in self.mapping.items()}
            mandy_df = final_raw.rename(columns=clean_mapping).reindex(columns=self.mandy_columns_list)

            # 清理文字與日期
            date_cols = ['Customer Reference Date', 'ETD', 'CRD']
            for c in date_cols: mandy_df[c] = mandy_df[c].apply(self.smart_date_format)
            other = [c for c in mandy_df.columns if c not in date_cols]
            mandy_df[other] = mandy_df[other].astype(str).replace(['nan', 'None', 'nan\n'], '')

            # 存檔
            ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            mandy_df.to_excel(f"mandy_output_{ts}.xlsx", index=False)
            self.append_to_mandy_sheet(mandy_df)
            
            # UI 顯示
            for i in self.tree.get_children(): self.tree.delete(i)
            self.tree["columns"] = self.tree_view_cols
            for c in self.tree_view_cols: self.tree.heading(c, text=c)
            for r in mandy_df[self.tree_view_cols].head(100).values.tolist(): self.tree.insert("", tk.END, values=r)
            
            self.info_label.config(text=f"Done! Processed {len(mandy_df)} rows.")
            messagebox.showinfo("Success", "Process Complete!")

        except Exception as e:
            messagebox.showerror("Error", str(e))

    def smart_date_format(self, x):
        if pd.isna(x) or str(x).strip().lower() in ['nan', 'none', '']: return ''
        try: return pd.to_datetime(x).strftime('%Y/%m/%d')
        except: return str(x).strip()

    def append_to_mandy_sheet(self, df):
        try:
            wb = openpyxl.load_workbook(self.output_target_path)
            ws = wb['Mandy'] if 'Mandy' in wb.sheetnames else wb.create_sheet('Mandy')
            if ws.max_row == 1 and ws.cell(1,1).value is None:
                for j, col in enumerate(df.columns, 1): ws.cell(row=1, column=j, value=col)
            start_r = ws.max_row + 1
            for i, row in enumerate(dataframe_to_rows(df, index=False, header=False), start_r):
                for j, val in enumerate(row, 1): ws.cell(row=i, column=j, value=val)
            wb.save(self.output_target_path)
            wb.close()
        except Exception as e: print(f"Append Error: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelProcessorApp(root)
    root.mainloop()