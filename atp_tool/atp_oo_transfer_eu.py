import pandas as pd
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import datetime
import os

class ExcelProcessorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel Data Integrator - EU Region (V3.1 Fix)")
        self.root.geometry("1400x900")
        
        self.input_file_path = ""
        self.output_target_path = ""
        self.all_sheets = []
        self.sheet_checkboxes = {} 
        
        # 44 個原始欄位
        self.column_names_list = [
            '上傳', 'Sales', 'Sales Document', 'Sales Document Item', 'Sold-to Abbreviation', 
            'Sold-to Party', 'Ship-to Party', 'Document Date', 'Customer Reference', 
            'Customer PO Item No', 'Material', 'Module Material', 'Customer Material', 
            'Order Quantity', 'Unit Price (ZPR0)', 'CRD', 'ETD', 'ETA', 'Ship-to Country', 
            'Original CRD', 'Remark', 'Delivery', 'Tracking No', 'Invoice or GUI No', 
            'Close', 'Is Sample (Y)', 'Not Allow Partial Delivery', 'Is Gift Item (ZFTO)', 
            'FOC Invoice Price (ZIPR)', 'Customer Price Group', 'Check Result', 'Quotation No', 
            'Ship-to Abbreviation', 'Domestic OR Foreign Sales', 'Incoterms', 'Order Remark', 
            'BU', 'Interface', 'Flash / DRAM', 'ETD Reason', 'ETD History', 
            'End Customer PN', '(訂單變更) 新增項次', '來源項次'
        ]
        
        self.mapping = {
            '上傳': 'Upload', 'Sales Document': 'Sales Document', 'Sales Document Item': 'Sales Document Item',
            'Sold-to Abbreviation': 'Sold-to Abbreviation', 'Sold-to Party': 'Sold-to Party', 'Ship-to Party': 'Ship-to Party',
            'Document Date': 'Customer Reference Date', 'Customer Reference': 'Customer Reference', 'Customer PO Item No': 'Customer PO Item No',
            'Material': 'Material', 'Module Material': 'Module Material', 'Customer Material': 'Customer Material',
            'Order Quantity': 'Order Quantity', 'Unit Price (ZPR0)': 'Unit Price', 'CRD': 'CRD', 'ETD': 'ETD', 'ETA': 'ETA',
            'Remark': 'Remark', 'Delivery': 'Price Check', 'Close': 'Close', 'Is Sample (Y)': 'Is Sample (Y)',
            'Is Gift Item (ZFTO)': 'Is Gift Item (ZFTO)', 'FOC Invoice Price (ZIPR)': 'FOC Invoice Price',
            'Ship-to Abbreviation': 'Ship-to Abbreviation', 'Incoterms': 'Incoterm (Header)', 'BU': 'BU',
            'Interface': 'Interface', 'Flash / DRAM': 'Flash / DRAM', 'ETD Reason': 'ETD Reason',
            'ETD History': 'ETD History', 'End Customer PN': 'End Customer PN', '(訂單變更) 新增項次': '訂單變更類型',
            '來源項次': '拆項來源項次 / 取消原因'
        }

        self.mandy_columns_list = [
            'Upload', 'Close', 'SO Type', 'Source SO No', 'Source SO Item No', 'Sales Document',
            'Sales Document Item', 'Sold-to Abbreviation', 'Ship-to Abbreviation', 'Customer Reference Date',
            'Sold-to Party', 'Ship-to Party', 'Customer Reference', 'Customer PO Item No', 'Material',
            'Module Material', 'Customer Material', 'Order Quantity', 'Unit Price', 'CRD', 'ETD', 'ETA',
            'Is Sample (Y)', 'Remark', 'Price Check', 'Purchase Order Type', 'Sample/Loan SFID', 'Is Gift Item (ZFTO)',
            'FOC Invoice Price', 'Masking Price', 'Masking Price Currency', 'End Customer PO#', 'End Customer PN',
            'Incoterm (Header)', 'Inco.Location (Header)', 'Incoterm (Item)', 'Inco.Location (Item)',
            'Special Procurement Application No', 'Titop DN', 'Tiptop DN Item', 'Unit Price Currency',
            'Original Unit Price', 'BU', 'Interface', 'Flash / DRAM', 'ETD Reason', 'Allocation Date',
            'ETD History', '訂單變更類型', '拆項來源項次 / 取消原因', '(ZDR) \nItem Description'
        ]

        self.tree_view_cols = ['Customer Reference Date', 'Sold-to Party', 'Sales Document', 'Material', 'Order Quantity', 'ETD', 'Price Check']
        self.setup_ui()

    def setup_ui(self):
        main_paned = ttk.Panedwindow(self.root, orient=tk.HORIZONTAL)
        main_paned.pack(fill=tk.BOTH, expand=True)
        left_frame = ttk.Frame(main_paned, padding=15)
        main_paned.add(left_frame, weight=1)

        ttk.Label(left_frame, text="Step 1: Input EU File", font=('Arial', 10, 'bold')).pack(anchor=tk.W)
        ttk.Button(left_frame, text="📁 Open Source Excel", command=self.open_file).pack(fill=tk.X, pady=5)
        self.input_label = ttk.Label(left_frame, text="No source", wraplength=350, foreground="gray")
        self.input_label.pack(pady=2, anchor=tk.W)

        ttk.Separator(left_frame, orient='horizontal').pack(fill='x', pady=10)
        ttk.Label(left_frame, text="Step 2: Output Target", font=('Arial', 10, 'bold')).pack(anchor=tk.W)
        ttk.Button(left_frame, text="🎯 Set Target File", command=self.set_output_target).pack(fill=tk.X, pady=5)
        self.output_label = ttk.Label(left_frame, text="Default: Same as Source", wraplength=350, foreground="gray")
        self.output_label.pack(pady=2, anchor=tk.W)

        ttk.Separator(left_frame, orient='horizontal').pack(fill='x', pady=10)
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", self.filter_sheets)
        ttk.Entry(left_frame, textvariable=self.search_var).pack(fill=tk.X, pady=5)

        list_container = ttk.Frame(left_frame, relief="sunken", borderwidth=1)
        list_container.pack(fill=tk.BOTH, expand=True, pady=5)
        self.canvas = tk.Canvas(list_container, highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(list_container, orient="vertical", command=self.canvas.yview)
        self.scroll_frame = ttk.Frame(self.canvas)
        self.scroll_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        ttk.Button(left_frame, text="🚀 START PROCESS", command=self.process_data).pack(fill=tk.X, ipady=10, pady=10)

        right_frame = ttk.Frame(main_paned, padding=15)
        main_paned.add(right_frame, weight=3)
        self.info_label = ttk.Label(right_frame, text="Ready", font=('Arial', 11, 'bold'))
        self.info_label.pack(anchor=tk.W, pady=(0, 10))
        self.tree = ttk.Treeview(right_frame, show="headings")
        self.tree.pack(fill=tk.BOTH, expand=True)

    def open_file(self):
        file = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file: return
        self.input_file_path = file
        self.output_target_path = file
        self.input_label.config(text=os.path.basename(file), foreground="blue")
        try:
            xl = pd.ExcelFile(file)
            self.all_sheets = xl.sheet_names
            for widget in self.scroll_frame.winfo_children(): widget.destroy()
            self.sheet_checkboxes = {}
            for sheet in self.all_sheets:
                var = tk.BooleanVar(value=False)
                chk = ttk.Checkbutton(self.scroll_frame, text=sheet, variable=var)
                self.sheet_checkboxes[sheet] = (var, chk)
            self.filter_sheets()
        except Exception as e: messagebox.showerror("Error", str(e))

    def filter_sheets(self, *args):
        s = self.search_var.get().strip().lower()
        for name, (var, chk) in self.sheet_checkboxes.items():
            if s in name.lower(): chk.pack(anchor=tk.W, pady=1, fill=tk.X)
            else: chk.pack_forget()

    def set_output_target(self):
        file = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if file:
            self.output_target_path = file
            self.output_label.config(text=os.path.basename(file), foreground="blue")

    def smart_date_format(self, x):
        s = str(x).strip()
        if s.lower() in ['nan', 'none', '', 'nat']: return ''
        try:
            dt = pd.to_datetime(x)
            return dt.strftime('%Y/%m/%d')
        except:
            return s # 如果是 Cancelled，直接回傳原字串

    def process_data(self):
        selected_names = [n for n, (v, c) in self.sheet_checkboxes.items() if v.get()]
        if not selected_names:
            messagebox.showwarning("Warning", "Select sheets first!")
            return

        std_list = [str(c).strip().replace('\n', ' ') for c in self.column_names_list]
        self.info_label.config(text="Processing...")
        self.root.update()
        processed_dfs = []
        
        try:
            for s_name in selected_names:
                # 讀取整張工作表，不限制行數
                df_src = pd.read_excel(self.input_file_path, sheet_name=s_name, dtype=str)
                df_src.columns = [str(c).strip().replace('\n', ' ') for c in df_src.columns]
                
                # 清除完全空白的行（避免讀到 Excel 底部幾萬行空資料）
                df_src = df_src.dropna(how='all')

                # 對齊欄位
                df_aligned = pd.DataFrame(columns=std_list)
                common_cols = [c for c in std_list if c in df_src.columns]
                df_aligned = pd.concat([df_aligned, df_src[common_cols]], axis=0, ignore_index=True)

                # --- ETD 篩選邏輯 (V3.1 核心修正) ---
                if 'ETD' in df_aligned.columns:
                    # 先把原本的 ETD 內容存一份備份，避免轉換過程遺失
                    etd_raw = df_aligned['ETD'].astype(str).str.strip()
                    
                    # 嘗試轉換日期
                    temp_date = pd.to_datetime(etd_raw, errors='coerce')
                    
                    # 保留條件：
                    # 1. 是日期且 >= 2025/1/1
                    # 2. 不是日期 (NaT) -> 包含 Cancelled, 文字, 或空白
                    mask = (temp_date >= datetime.datetime(2025, 1, 1)) | (temp_date.isna())
                    df_aligned = df_aligned[mask].copy()

                processed_dfs.append(df_aligned)

            if not processed_dfs: return
            final_raw = pd.concat(processed_dfs, ignore_index=True)
            
            # Mapping
            clean_mapping = {str(k).strip().replace('\n', ' '): v for k, v in self.mapping.items()}
            mandy_df = final_raw.rename(columns=clean_mapping).reindex(columns=self.mandy_columns_list)

            # 格式化日期 (確保 Cancelled 被保留)
            date_cols = ['Customer Reference Date', 'ETD', 'CRD', 'ETA']
            for c in date_cols:
                if c in mandy_df.columns:
                    mandy_df[c] = mandy_df[c].apply(self.smart_date_format)
            
            # 其他欄位清理
            other_cols = [c for c in mandy_df.columns if c not in date_cols]
            mandy_df[other_cols] = mandy_df[other_cols].astype(str).replace(['nan', 'None', 'nan\n'], '')

            # 存檔與追加
            ts = datetime.datetime.now().strftime("%m%d_%H%M")
            mandy_df.to_excel(f"EU_Result_{ts}.xlsx", index=False)
            self.append_to_mandy_sheet(mandy_df)
            
            # 預覽更新
            for i in self.tree.get_children(): self.tree.delete(i)
            self.tree["columns"] = self.tree_view_cols
            for c in self.tree_view_cols: self.tree.heading(c, text=c)
            for r in mandy_df[self.tree_view_cols].head(50).values.tolist(): self.tree.insert("", tk.END, values=r)
            
            self.info_label.config(text=f"Done! Saved {len(mandy_df)} rows.")
            messagebox.showinfo("Success", f"Total {len(mandy_df)} rows processed.")

        except Exception as e:
            messagebox.showerror("Error", str(e))

    def append_to_mandy_sheet(self, df):
        try:
            wb = openpyxl.load_workbook(self.output_target_path)
            ws = wb['Mandy'] if 'Mandy' in wb.sheetnames else wb.create_sheet('Mandy')
            start_r = ws.max_row + 1
            if ws.max_row == 1 and ws.cell(1,1).value is None:
                for j, col in enumerate(df.columns, 1): ws.cell(row=1, column=j, value=col)
                start_r = 2
            for i, row in enumerate(dataframe_to_rows(df, index=False, header=False), start_r):
                for j, val in enumerate(row, 1): ws.cell(row=i, column=j, value=val)
            wb.save(self.output_target_path)
        except Exception as e: print(f"Append Error: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelProcessorApp(root)
    root.mainloop()